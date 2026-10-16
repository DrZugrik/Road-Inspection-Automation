import unittest

from docx import Document
from docx.shared import Pt

from docx_placeholders import replace_in_document, replace_in_paragraph


class DocxPlaceholderTests(unittest.TestCase):
    def test_replaces_placeholder_split_between_runs_and_keeps_surrounding_styles(self):
        document = Document()
        paragraph = document.add_paragraph()
        prefix = paragraph.add_run('Дорога: ')
        prefix.bold = True
        start = paragraph.add_run('{{ДОРОГА_')
        start.italic = True
        end = paragraph.add_run('НАЗВАНИЕ}}')
        suffix = paragraph.add_run(' готово')
        suffix.font.size = Pt(14)

        count = replace_in_paragraph(paragraph, '{{ДОРОГА_НАЗВАНИЕ}}', 'М-10')

        self.assertEqual(count, 1)
        self.assertEqual(paragraph.text, 'Дорога: М-10 готово')
        self.assertTrue(prefix.bold)
        self.assertTrue(start.italic)
        self.assertEqual(suffix.font.size, Pt(14))

    def test_replaces_multiple_occurrences_without_recreating_paragraph(self):
        document = Document()
        paragraph = document.add_paragraph('{{X}} и {{X}}')
        self.assertEqual(replace_in_paragraph(paragraph, '{{X}}', 'значение'), 2)
        self.assertEqual(paragraph.text, 'значение и значение')

    def test_replaces_placeholder_in_table_cell(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        paragraph = cell.paragraphs[0]
        paragraph.add_run('ID: ')
        paragraph.add_run('{{ДОРОГА_')
        paragraph.add_run('ИД}}')

        count = replace_in_document(document, {'{{ДОРОГА_ИД}}': '77-ОП-РЗ-77К-001'})

        self.assertEqual(count, 1)
        self.assertEqual(cell.text, 'ID: 77-ОП-РЗ-77К-001')

    def test_replaces_placeholder_in_header(self):
        document = Document()
        header = document.sections[0].header
        header.paragraphs[0].add_run('{{НОМЕР_ПП}}')

        self.assertEqual(replace_in_document(document, {'{{НОМЕР_ПП}}': 12}), 1)
        self.assertEqual(header.paragraphs[0].text, '12')

    def test_missing_placeholder_returns_zero(self):
        document = Document()
        paragraph = document.add_paragraph('Обычный текст')
        self.assertEqual(replace_in_paragraph(paragraph, '{{НЕ_СУЩЕСТВУЕТ}}', 'x'), 0)

    def test_can_remove_split_placeholder_without_losing_other_runs(self):
        document = Document()
        paragraph = document.add_paragraph()
        prefix = paragraph.add_run('До ')
        prefix.bold = True
        paragraph.add_run('{{ФОТО_')
        paragraph.add_run('ВИД}}')
        suffix = paragraph.add_run(' после')
        suffix.italic = True

        self.assertEqual(replace_in_document(document, {'{{ФОТО_ВИД}}': ''}), 1)
        self.assertEqual(paragraph.text, 'До  после')
        self.assertTrue(prefix.bold)
        self.assertTrue(suffix.italic)